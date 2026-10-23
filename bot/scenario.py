"""Загрузка и проверка сценария из Excel / Google Таблицы.

Сценарий целиком живёт в таблице с листами:
  Настройки  - общие параметры и служебные тексты
  Шаги       - сообщения бота (вопросы, информационные блоки)
  Кнопки     - варианты ответов для каждого шага
  Маршруты   - порядок шагов в каждой ветке диалога
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

import openpyxl

# Типы шагов (значения колонки "Тип" на листе "Шаги")
CHOICE = "кнопки"        # один вариант из списка
MULTI = "несколько"      # несколько вариантов + кнопка "Готово"
MENU = "меню"            # список разделов, в каждый можно зайти и вернуться
TEXT = "текст"           # свободный ввод текста (можно добавить и кнопки)
PHONE = "телефон"        # ввод номера телефона или кнопка "Отправить номер"
FILE = "файл"            # клиент прикрепляет файл
MESSAGE = "сообщение"    # просто текст, бот сразу идёт дальше
FINAL = "финал"          # завершение: заявка уходит менеджеру

STEP_TYPES = (CHOICE, MULTI, MENU, TEXT, PHONE, FILE, MESSAGE, FINAL)
WITH_OPTIONS = (CHOICE, MULTI, MENU, FINAL)

# Особое значение перехода: вернуться на предыдущий шаг (например, из раздела в меню разделов)
BACK = "НАЗАД"

CHECKS = ("", "ссылка", "email")

# Цвет кнопки (колонка "Цвет кнопки") → значение style в Telegram
COLORS = {"": "", "зелёная": "success", "зеленая": "success", "синяя": "primary", "красная": "danger"}

DEFAULT_SETTINGS = {
    "Стартовый шаг": "приватность",
    "Главное меню": "главное_меню",
    "Чат для заявок (ID)": "",
    "Администраторы (ID через запятую)": "",
    "Напоминание через, минут (0 = выключено)": "60",
    "Текст напоминания": "Вы не закончили заявку. Осталось совсем немного, продолжим?",
    'Кнопка "Продолжить" в напоминании': "Продолжить",
    'Кнопка "Назад"': "Назад",
    'Кнопка "Готово" по умолчанию': "Готово",
    'Цвет кнопки "Готово"': "зелёная",
    'Кнопка "Отправить номер"': "Отправить мой номер",
    'Кнопка "В главное меню"': "В главное меню",
    'Сообщение "выберите кнопку"': "Пожалуйста, выберите вариант с помощью кнопок выше.",
    'Сообщение "выберите хотя бы один"': "Выберите хотя бы один вариант",
    'Сообщение "неверный телефон"': "Не получилось распознать номер. Напишите его в формате +7 999 123-45-67",
    'Сообщение "неверная ссылка"': "Похоже, это не ссылка. Пришлите адрес сайта, например site.ru",
    'Сообщение "неверный email"': "Похоже, это не email. Пример: name@mail.ru",
    'Сообщение "нужен файл"': "Прикрепите, пожалуйста, файл (скрепка внизу экрана)",
    'Сообщение "файл не нужен"': "Сейчас файл не нужен, выберите вариант с помощью кнопок.",
    "Ответ на вопрос клиента": "Спасибо! Передали ваш вопрос специалисту, ответ придёт сюда.",
    "Подпись к ответу специалиста": "Ответ специалиста:",
    "Обновлять сценарий каждые, минут": "5",
}


class ScenarioError(Exception):
    def __init__(self, errors: list[str]):
        super().__init__("\n".join(errors))
        self.errors = errors


@dataclass
class Option:
    text: str
    goto: str = ""
    prompt: str = ""
    exclusive: bool = False
    url: str = ""
    color: str = ""
    icon: str = ""
    row: int = 0


@dataclass
class Step:
    id: str
    type: str
    text: str
    label: str = ""
    next: str = ""
    back: bool = True
    images: list[str] = field(default_factory=list)
    check: str = ""
    done_label: str = ""
    notify: bool = False
    options: list[Option] = field(default_factory=list)
    row: int = 0


@dataclass
class Scenario:
    steps: dict[str, Step]
    routes: dict[str, list[str]]
    settings: dict[str, str]

    def setting(self, key: str) -> str:
        value = self.settings.get(key, "").strip()
        return value or DEFAULT_SETTINGS.get(key, "")

    def setting_int(self, key: str) -> int:
        try:
            return int(float(self.setting(key).replace(",", ".")))
        except ValueError:
            return int(DEFAULT_SETTINGS.get(key, "0") or 0)

    def setting_ids(self, key: str) -> list[int]:
        ids = []
        for part in re.split(r"[,\s;]+", self.setting(key)):
            if re.fullmatch(r"-?\d+", part):
                ids.append(int(part))
        return ids

    @property
    def start(self) -> str:
        return self.setting("Стартовый шаг")

    @property
    def main(self) -> str:
        return self.setting("Главное меню")


# ---------- чтение таблицы ----------

def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).replace("\r\n", "\n").strip()


def _yes(value: str) -> bool:
    return value.strip().lower() in ("да", "yes", "+", "1", "true", "истина")


def _sheet_rows(wb, name: str, columns: dict[str, str], errors: list[str]):
    """Строки листа как словари {ключ: значение}; колонки ищутся по заголовку."""
    if name not in wb.sheetnames:
        errors.append(f"Нет листа \"{name}\"")
        return []
    ws = wb[name]
    rows = ws.iter_rows(values_only=True)
    header = [_cell(h).lower() for h in next(rows, [])]
    index = {}
    for key, title in columns.items():
        found = [i for i, h in enumerate(header) if h.startswith(title.lower())]
        if found:
            index[key] = found[0]
        elif not key.startswith("_"):
            errors.append(f"Лист \"{name}\": нет колонки \"{title}\"")
    result = []
    for n, row in enumerate(rows, start=2):
        item = {key: _cell(row[i]) if i < len(row) else "" for key, i in index.items()}
        if any(item.values()):
            item["row"] = n
            result.append(item)
    return result


STEP_COLUMNS = {
    "id": "ID шага",
    "type": "Тип",
    "text": "Текст сообщения",
    "label": "Название в заявке",
    "next": "Следующий шаг",
    "back": 'Кнопка "Назад"',
    "_images": "Картинки",
    "_check": "Проверка ввода",
    "_done": 'Текст кнопки "Готово"',
    "_notify": "Сразу сообщить менеджеру",
}

OPTION_COLUMNS = {
    "step": "ID шага",
    "text": "Текст кнопки",
    "goto": "Переход",
    "prompt": "Попросить ввести",
    "_exclusive": "Снимает другие",
    "_url": "Ссылка",
    "_color": "Цвет кнопки",
    "_icon": "Иконка",
}


def parse_workbook(data: bytes | str) -> Scenario:
    """Читает xlsx (путь или байты), проверяет и возвращает сценарий.
    При ошибках бросает ScenarioError со списком понятных сообщений."""
    source = io.BytesIO(data) if isinstance(data, bytes) else data
    try:
        wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    except Exception as e:  # битый файл, не тот формат и т.п.
        raise ScenarioError([f"Не удалось открыть таблицу: {e}"])

    errors: list[str] = []

    settings = {}
    if "Настройки" in wb.sheetnames:
        for row in wb["Настройки"].iter_rows(min_row=2, values_only=True):
            if row and _cell(row[0]):
                settings[_cell(row[0])] = _cell(row[1]) if len(row) > 1 else ""

    steps: dict[str, Step] = {}
    for r in _sheet_rows(wb, "Шаги", STEP_COLUMNS, errors):
        sid, stype = r.get("id", ""), r.get("type", "").lower()
        where = f"Шаги, строка {r['row']}"
        if not sid:
            errors.append(f"{where}: не заполнен ID шага")
            continue
        if sid in steps:
            errors.append(f"{where}: ID \"{sid}\" уже есть в строке {steps[sid].row}")
            continue
        if stype not in STEP_TYPES:
            errors.append(f"{where}: неизвестный тип \"{r.get('type')}\". Допустимые: {', '.join(STEP_TYPES)}")
            continue
        check = r.get("_check", "").lower()
        if check not in CHECKS:
            errors.append(f"{where}: неизвестная проверка \"{check}\". Допустимые: ссылка, email или пусто")
            check = ""
        steps[sid] = Step(
            id=sid,
            type=stype,
            text=r.get("text", ""),
            label=r.get("label", ""),
            next=r.get("next", ""),
            back=not r.get("back") or _yes(r.get("back", "")),
            images=[u for u in re.split(r"\s+", r.get("_images", "")) if u],
            check=check,
            done_label=r.get("_done", ""),
            notify=_yes(r.get("_notify", "")),
            row=r["row"],
        )
        if not steps[sid].text:
            errors.append(f"{where}: пустой текст сообщения у шага \"{sid}\"")

    for r in _sheet_rows(wb, "Кнопки", OPTION_COLUMNS, errors):
        where = f"Кнопки, строка {r['row']}"
        step = steps.get(r.get("step", ""))
        if not step:
            errors.append(f"{where}: шага \"{r.get('step')}\" нет на листе \"Шаги\"")
            continue
        if not r.get("text"):
            errors.append(f"{where}: пустой текст кнопки")
            continue
        step.options.append(Option(
            text=r["text"],
            goto=r.get("goto", ""),
            prompt=r.get("prompt", ""),
            exclusive=_yes(r.get("_exclusive", "")),
            url=r.get("_url", ""),
            color=r.get("_color", "").lower(),
            icon=r.get("_icon", ""),
            row=r["row"],
        ))

    routes: dict[str, list[str]] = {}
    if "Маршруты" in wb.sheetnames:
        for n, row in enumerate(wb["Маршруты"].iter_rows(min_row=2, values_only=True), start=2):
            name = _cell(row[0]) if row else ""
            if not name:
                continue
            items = [x.strip() for x in re.split(r"[,\n;]+", _cell(row[1]) if len(row) > 1 else "") if x.strip()]
            if name in steps:
                errors.append(f"Маршруты, строка {n}: имя ветки \"{name}\" совпадает с ID шага")
            elif name in routes:
                errors.append(f"Маршруты, строка {n}: ветка \"{name}\" уже есть")
            elif not items:
                errors.append(f"Маршруты, строка {n}: у ветки \"{name}\" не указаны шаги")
            else:
                routes[name] = items
    else:
        errors.append('Нет листа "Маршруты"')

    scenario = Scenario(steps=steps, routes=routes, settings=settings)
    errors += validate(scenario)
    if errors:
        raise ScenarioError(errors)
    return scenario


def validate(sc: Scenario) -> list[str]:
    errors = []

    def known(target: str) -> bool:
        return target in sc.steps or target in sc.routes or target == BACK

    for key in ("Стартовый шаг", "Главное меню"):
        if sc.setting(key) not in sc.steps:
            errors.append(f"Настройки: \"{key}\" = \"{sc.setting(key)}\", такого шага нет")

    for step in sc.steps.values():
        where = f"Шаги, строка {step.row} (\"{step.id}\")"
        if step.next and not known(step.next):
            errors.append(f"{where}: следующий шаг \"{step.next}\" не найден")
        if step.type in WITH_OPTIONS and not step.options:
            errors.append(f"{where}: для типа \"{step.type}\" нужны кнопки на листе \"Кнопки\"")
        for opt in step.options:
            if opt.goto and not known(opt.goto):
                errors.append(f"Кнопки, строка {opt.row}: переход \"{opt.goto}\" не найден ни среди шагов, ни среди веток")
            if opt.url and not re.match(r"https?://", opt.url):
                errors.append(f"Кнопки, строка {opt.row}: ссылка должна начинаться с http:// или https://")
            if opt.url and (opt.goto or opt.prompt):
                errors.append(f"Кнопки, строка {opt.row}: кнопка-ссылка не может иметь переход или поле ввода")
            if opt.color not in COLORS:
                errors.append(f"Кнопки, строка {opt.row}: цвет \"{opt.color}\" не поддерживается. "
                              "Допустимые: зелёная, синяя, красная или пусто")
            if opt.icon and not opt.icon.isdigit():
                errors.append(f"Кнопки, строка {opt.row}: в колонке \"Иконка\" нужен числовой ID кастомного эмодзи")
        for url in step.images:
            if not re.match(r"https?://", url):
                errors.append(f"{where}: картинка \"{url}\" должна быть ссылкой http(s)://")

    for name, items in sc.routes.items():
        for item in items:
            if item not in sc.steps and item not in sc.routes:
                errors.append(f"Маршруты, ветка \"{name}\": шаг \"{item}\" не найден")
            elif item in sc.steps and sc.steps[item].type == MESSAGE and item == items[-1]:
                errors.append(f"Маршруты, ветка \"{name}\": ветка не может заканчиваться сообщением")
    return errors


# ---------- источники ----------

GOOGLE_RE = re.compile(r"docs\.google\.com/spreadsheets/d/([\w-]+)")


def export_url(source: str) -> str | None:
    """Ссылка на выгрузку Google Таблицы в xlsx (таблица должна быть открыта по ссылке на просмотр)."""
    m = GOOGLE_RE.search(source)
    return f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=xlsx" if m else None


async def load(source: str) -> Scenario:
    url = export_url(source)
    if not url:
        return parse_workbook(source)
    import aiohttp
    try:
        async with aiohttp.ClientSession(trust_env=True) as http:
            async with http.get(url, timeout=aiohttp.ClientTimeout(total=30)) as resp:
                data = await resp.read()
                if resp.status != 200 or not data.startswith(b"PK"):
                    raise ScenarioError([
                        f"Google Таблица недоступна (код {resp.status}). "
                        'Проверьте доступ: "Все, у кого есть ссылка" → "Читатель"'])
    except aiohttp.ClientError as e:
        raise ScenarioError([f"Не удалось скачать Google Таблицу: {e}"])
    return parse_workbook(data)
